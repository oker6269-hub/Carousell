#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""用当前的标题清洗规则，重新清洗已有表格里的 ProductNameCn。

场景：规则改过之后，历史行还留着旧结果（例如 全新正貨… 被清成了 正貨…）。

默认只预览（dry-run），加 --apply 才写入，且写入前会备份为
`素材保存表格.xlsx.bak-<时间戳>`，出问题直接改名回去即可。

用法：
    python tools/fix_titles.py /Users/max/carousell/新品采集10_25
    python tools/fix_titles.py /Users/max/carousell/新品采集10_25 --apply
"""
import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook          # noqa: E402

from app.storage import HEADERS, XLSX_NAME  # noqa: E402
from app.title_rules import TitleRules      # noqa: E402

TITLE_COL = HEADERS.index("ProductNameCn") + 1
SKU_COL = HEADERS.index("SKU") + 1


def main() -> int:
    ap = argparse.ArgumentParser(description="重新清洗已有表格的 ProductNameCn")
    ap.add_argument("root", help="素材根目录（包含 素材保存表格.xlsx）")
    ap.add_argument("--apply", action="store_true", help="真的写入（默认只预览）")
    ap.add_argument("--limit", type=int, default=15, help="预览显示条数（默认 15）")
    args = ap.parse_args()

    path = Path(args.root) / XLSX_NAME
    if not path.exists():
        print(f"找不到表格: {path}")
        return 1

    rules = TitleRules(log=print)
    if not rules.enabled:
        print("标题清洗当前是关闭状态，先启用再跑。")
        return 1

    wb = load_workbook(path)
    ws = wb["Sheet1"]
    changes = []
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=TITLE_COL)
        old = str(cell.value or "")
        if not old:
            continue
        new = rules.clean(old)
        if new and new != old:
            changes.append((row, ws.cell(row=row, column=SKU_COL).value, old, new))

    print(f"表格: {path}")
    print(f"共 {ws.max_row - 1} 行数据，需要改写 {len(changes)} 行\n")
    for row, sku, old, new in changes[: args.limit]:
        print(f"  行{row} [{sku}]")
        print(f"    - {old[:70]}")
        print(f"    + {new[:70]}")
    if len(changes) > args.limit:
        print(f"  … 其余 {len(changes) - args.limit} 行略")

    if not changes:
        return 0
    if not args.apply:
        print("\n（预览模式）确认无误后加 --apply 写入，会先生成 .bak 备份。")
        return 0

    backup = path.with_name(f"{path.name}.bak-{datetime.now():%Y%m%d-%H%M%S}")
    shutil.copy2(path, backup)
    for row, _sku, _old, new in changes:
        ws.cell(row=row, column=TITLE_COL).value = new
    try:
        wb.save(path)
    except PermissionError:
        print(f"\n写入失败：{path.name} 正被 Excel 打开，请关闭后重试（备份已生成：{backup.name}）")
        return 2
    print(f"\n已写入 {len(changes)} 行；备份：{backup.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
